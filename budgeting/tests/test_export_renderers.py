"""Testes dos arquivos reais. Não requer banco de dados."""
import json
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from unittest import TestCase
from zipfile import ZipFile

import openpyxl
from pypdf import PdfReader
from budgeting.services.budget_export_renderers import render_xlsx, render_pdf, render_pdf_bundle
from budgeting.services.budget_export_branding import COMPANY, work_location


class ExportRenderersTests(TestCase):
    def setUp(self):
        self.data = json.loads((Path(__file__).parent / 'fixtures/export_pc25017.json').read_text())

    def test_client_excel_has_only_six_columns_and_no_internal_content(self):
        data = deepcopy(self.data)
        data['audience'] = 'CLIENT'
        # Mesmo se um chamador enviar dados internos, o renderizador não os utiliza.
        data['stages'][0]['items'][0]['cost']['components'][0]['description_snapshot'] = 'SEGREDO_INTERNO'
        data['stages'][0]['items'][0]['description'] = '=HYPERLINK("https://example.invalid", "texto")'
        book = openpyxl.load_workbook(BytesIO(render_xlsx(data)), data_only=False)
        self.assertEqual(book.sheetnames, ['Planilha'])
        sheet = book.active
        self.assertEqual(sheet.max_column, 6)
        self.assertFalse(any(c.data_type == 'f' for row in sheet for c in row))
        self.assertFalse(any(c.hyperlink for row in sheet for c in row))
        text = ' '.join(str(c.value) for row in sheet for c in row)
        for term in ('SEGREDO_INTERNO', 'BDI', 'DESPESAS INDIRETAS', '13110.4775'):
            self.assertNotIn(term, text)
        self.assertEqual(sheet['F11'].value, 22000)  # total da etapa na própria linha da etapa
        self.assertEqual(sheet['F16'].value, 22000)  # total geral

    def test_internal_excel_has_formulas_and_correct_cost_and_sale_totals(self):
        output = render_xlsx(self.data)
        formulas = openpyxl.load_workbook(BytesIO(output), data_only=False)
        values = openpyxl.load_workbook(BytesIO(output), data_only=True)
        self.assertEqual(set(formulas.sheetnames), {'Planilha', 'DI', 'DI Auxiliares', 'BDI', 'Composições'})
        self.assertEqual(values['Planilha']['F11'].value, 22000)
        self.assertEqual(values['Planilha']['F16'].value, 22000)
        self.assertEqual(values['Planilha']['N16'].value, 16800)
        self.assertEqual(values['Planilha']['S12'].value, 13879.79)
        self.assertEqual(values['Planilha']['T12'].value, 0)
        self.assertAlmostEqual(values['BDI']['B17'].value, 1.309712613574, places=12)
        self.assertAlmostEqual(values['BDI']['B12'].value, 2445.081175, places=8)
        self.assertEqual(formulas['Planilha']['P12'].value, '=IF(O12="Sim",\'BDI\'!$B$17,1)')
        self.assertEqual(formulas['Planilha']['Q12'].value, '=IF(O12="Sim",ROUND(J12*P12,2),J12)')
        self.assertIn('DI_Material', formulas.defined_names)
        # Agrupador/desativado não repete os custos da folha ativa.
        self.assertEqual(values['Composições']['L11'].value, 0)
        self.assertEqual(values['Composições']['L12'].value, 10600)
        self.assertEqual(values['Composições']['L13'].value, 0)

    def test_client_pdf_does_not_expose_memory_or_snapshot(self):
        data = deepcopy(self.data); data['audience'] = 'CLIENT'
        pdf = PdfReader(BytesIO(render_pdf(data)))
        text = '\n'.join(page.extract_text() for page in pdf.pages)
        self.assertIn('22.000,00', text)
        self.assertIn('13.879,79', text)
        for term in ('BDI', 'D.I.', 'INSUMO_INTERNO_DESATIVADO', '10.600,00', '13110'):
            self.assertNotIn(term, text)

    def test_internal_bundle_has_three_separate_landscape_reports(self):
        with ZipFile(BytesIO(render_pdf_bundle(self.data))) as archive:
            self.assertEqual(archive.namelist(), [f'PC25017-R00-DARUIX-{kind}.pdf' for kind in ('PLANILHA', 'BDI', 'DI')])
            self.assertIsNone(archive.testzip())
            texts = {}
            for name in archive.namelist():
                pdf = PdfReader(BytesIO(archive.read(name)))
                texts[name.split('-')[-1][:-4]] = '\n'.join(page.extract_text() for page in pdf.pages)
                self.assertIn('Página 1', pdf.pages[0].extract_text())
                for page in pdf.pages:
                    self.assertGreater(float(page.mediabox.width), float(page.mediabox.height))
                    self.assertEqual(page.get('/Rotate', 0) % 360, 0)
            for term in ('22.000,00', 'CUSTOS E FORMAÇÃO', 'INSUMO_INTERNO_DESATIVADO'):
                self.assertIn(term, texts['PLANILHA'])
            self.assertNotIn('MEMÓRIA DE BDI', texts['PLANILHA'])
            self.assertNotIn('DESPESAS INDIRETAS', texts['PLANILHA'])
            self.assertIn('22.003,17', texts['BDI'])
            self.assertIn('MEMÓRIA DE BDI', texts['BDI'])
            self.assertNotIn('INSUMO_INTERNO_DESATIVADO', texts['BDI'])
            self.assertNotIn('13.110,4775', texts['BDI'])
            self.assertIn('DESPESAS INDIRETAS', texts['DI'])
            self.assertIn('13.110,4775', texts['DI'])
            self.assertIn('Alimentação pessoal direto', texts['DI'])
            self.assertNotIn('INSUMO_INTERNO_DESATIVADO', texts['DI'])

    def test_client_bundle_has_three_reports_without_calculation_details(self):
        data = deepcopy(self.data); data['audience'] = 'CLIENT'
        # O renderizador também filtra, mesmo se receber a memória completa.
        with ZipFile(BytesIO(render_pdf_bundle(data))) as archive:
            self.assertEqual(archive.namelist(), [f'PC25017-R00-CLIENTE-{kind}.pdf' for kind in ('PLANILHA', 'BDI', 'DI')])
            texts = {}
            for name in archive.namelist():
                pdf = PdfReader(BytesIO(archive.read(name)))
                text = '\n'.join(page.extract_text() for page in pdf.pages)
                texts[name.split('-')[-1][:-4]] = text
                for term in ('ARRED(', 'CUSTOS E FORMAÇÃO', 'MEMÓRIA DE BDI', 'Fator BDI',
                             'INSUMO_INTERNO_DESATIVADO', 'Base sem arredondamento', '13.110,4775',
                             'Coeficientes', 'USO INTERNO', 'D / (1 -', 'Lucro / administração'):
                    self.assertNotIn(term, text)
                for page in pdf.pages:
                    self.assertGreater(float(page.mediabox.width), float(page.mediabox.height))
                    self.assertEqual(page.get('/Rotate', 0) % 360, 0)
            self.assertIn('22.000,00', texts['PLANILHA'])
            self.assertNotIn('BDI', texts['PLANILHA'])
            self.assertIn('30,9713%', texts['BDI'])
            self.assertIn('2.445,08', texts['DI'])
            self.assertNotIn('Mão de obra', texts['DI'])

    def test_pdf_rejects_unknown_audience_and_document(self):
        with self.assertRaises(ValueError):
            render_pdf(self.data, 'ALL')
        with self.assertRaises(ValueError):
            render_pdf({**self.data, 'audience': 'UNKNOWN'})

    def test_bundle_filenames_cannot_create_paths(self):
        data = deepcopy(self.data); data['proposal_code'] = '../../PC/25017'
        with ZipFile(BytesIO(render_pdf_bundle(data))) as archive:
            self.assertEqual(len(archive.namelist()), 3)
            for name in archive.namelist():
                self.assertNotIn('/', name)
                self.assertNotIn('\\', name)

    def test_long_client_budget_paginates_and_repeats_headers(self):
        data = deepcopy(self.data); data['audience'] = 'CLIENT'
        first = data['stages'][0]['items'][0]
        data['stages'][0]['items'] = [dict(first, number=f'1.{i}',description='Serviço com descrição longa. '*18) for i in range(1,45)]
        pdf = PdfReader(BytesIO(render_pdf(data)))
        self.assertGreater(len(pdf.pages), 1)
        for page in pdf.pages:
            self.assertIn('Discriminação', page.extract_text())
            for text in ('PC25017', 'Cliente:', 'Local:', 'Serviços:', 'Empresa:', 'REVISÃO Nº'):
                self.assertIn(text, page.extract_text())
            self.assertGreaterEqual(len(page.images), 1)

    def test_all_excel_sheets_have_branding_and_print_titles(self):
        data = deepcopy(self.data)
        data['work_location'] = 'Avenida de teste, 123 - São Paulo/SP'
        book = openpyxl.load_workbook(BytesIO(render_xlsx(data)))
        for sheet in book:
            row3 = ' '.join(str(cell.value) for cell in sheet[3] if cell.value is not None)
            row7 = ' '.join(str(cell.value) for cell in sheet[7] if cell.value is not None)
            self.assertIn('Local: ' + data['work_location'], row3)
            self.assertIn('REVISÃO Nº 00', row7)
            self.assertEqual(len(sheet._images), 1)
            self.assertEqual(sheet.print_title_rows, '$1:$10')
            self.assertEqual(sheet.freeze_panes, 'C11')
            self.assertTrue(sheet['A9'].value)
            self.assertIn(COMPANY, sheet.oddFooter.left.text)
            self.assertEqual(sheet.oddFooter.right.text, '&G')

    def test_bdi_di_use_original_palette(self):
        book = openpyxl.load_workbook(BytesIO(render_xlsx(self.data)))
        def fill(sheet, cell):
            return book[sheet][cell].fill.fgColor.rgb[-6:]
        self.assertEqual(fill('BDI', 'A1'), 'FFFFCC')
        self.assertEqual(fill('BDI', 'B17'), 'FFFFCC')
        self.assertEqual(fill('BDI', 'B19'), 'CCFFCC')
        self.assertEqual(fill('BDI', 'B23'), 'CCFFCC')
        self.assertEqual(fill('BDI', 'A24'), '99FF33')
        self.assertEqual(fill('DI', 'A1'), 'CCFFFF')
        self.assertEqual(fill('DI', 'A11'), 'CCFFCC')
        self.assertEqual(fill('DI', 'D12'), 'FFFFCC')
        self.assertEqual(fill('DI', 'E12'), 'CCFFFF')
        rows = {row[0].value: row[0].row for row in book['DI'] if row[0].value}
        self.assertEqual(fill('DI', f"B{rows['01.09']}"), 'FFDDFF')
        self.assertEqual(fill('DI', f"F{rows['04.01']}"), '99FF33')

    def test_work_location_uses_project_data_and_handles_missing_fields(self):
        from types import SimpleNamespace
        self.assertEqual(work_location(None, None), '')
        self.assertEqual(work_location(SimpleNamespace(work_address='Rua B'), None), 'Rua B')
        self.assertEqual(work_location(SimpleNamespace(work_address='Rua B'),
                                       SimpleNamespace(address='Rua A')), 'Rua A')
        self.assertEqual(work_location(SimpleNamespace(address='Endereço do cliente'), None), '')
